import json
import sys
import openpyxl


def leaf_data(row):
    data = {"ipa": row.get("IPA")}
    for key, header in (("orthography", "ORTHOGRAPHY"), ("pos", "POS"), ("eng", "ENG"),
                         ("fra", "FRA"), ("cf", "CF"), ("var", "VAR"), ("other", "OTHER"),
                         ("notes", "NOTES")):
        val = row.get(header)
        if val is not None and val != "":
            data[key] = val
    return data


def convert(rows):
    words = []
    word = sense = subdef = None
    example_target = None  # dict that "examples" get appended to (a subdef, derived, or related)
    warnings = []

    def new_word(row):
        nonlocal word, sense, subdef, example_target
        word = {"ipa": row.get("IPA"), "senses": []}
        if row.get("ORTHOGRAPHY"):
            word["orthography"] = row.get("ORTHOGRAPHY")
        words.append(word)
        new_sense(row)

    def new_sense(row):
        nonlocal sense, subdef, example_target
        sense = {"subdefinitions": []}
        word["senses"].append(sense)
        new_subdef(row)

    def new_subdef(row):
        nonlocal subdef, example_target
        subdef = leaf_data(row)
        sense["subdefinitions"].append(subdef)
        example_target = subdef

    for row in rows:
        rtype = row.get("TYPE")
        index = row.get("INDEX")
        if rtype is None or rtype == "":
            warnings.append(f"row {index}: blank TYPE, treated as 'word'")
            rtype = "word"

        if rtype == "word":
            new_word(row)
        elif rtype == "sense":
            if word is None:
                warnings.append(f"row {index}: 'sense' with no open word, treated as 'word'")
                new_word(row)
            else:
                new_sense(row)
        elif rtype == "subdef":
            if sense is None:
                warnings.append(f"row {index}: 'subdef' with no open sense, treated as 'word'")
                new_word(row)
            else:
                new_subdef(row)
        elif rtype in ("derived", "related"):
            if subdef is None:
                warnings.append(f"row {index}: '{rtype}' with no open subdefinition, treated as 'word'")
                new_word(row)
            else:
                d = leaf_data(row)
                subdef.setdefault(rtype, []).append(d)
                example_target = d
        elif rtype == "example":
            if example_target is None:
                warnings.append(f"row {index}: 'example' with no open subdefinition/derived/related, treated as 'word'")
                new_word(row)
            else:
                ex = {"ipa": row.get("IPA")}
                if row.get("ORTHOGRAPHY"):
                    ex["orthography"] = row.get("ORTHOGRAPHY")
                if row.get("ENG"):
                    ex["eng"] = row.get("ENG")
                if row.get("FRA"):
                    ex["fra"] = row.get("FRA")
                if row.get("NOTES"):
                    ex["notes"] = row.get("NOTES")
                example_target.setdefault("examples", []).append(ex)
        else:
            warnings.append(f"row {index}: unknown TYPE '{rtype}', treated as 'word'")
            new_word(row)

    return words, warnings


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "obolo_dict.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else src.rsplit(".", 1)[0] + ".json"

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    raw_rows = list(ws.iter_rows(min_row=2, values_only=True))
    rows = [dict(zip(headers, r)) for r in raw_rows if any(v is not None for v in r)]

    words, warnings = convert(rows)

    with open(out, "w", encoding="utf-8") as f:
        json.dump(words, f, ensure_ascii=False, indent=2)

    print(f"Wrote {len(words)} word entries to {out}")
    if warnings:
        print(f"\n{len(warnings)} warning(s):", file=sys.stderr)
        for w in warnings:
            print(" -", w, file=sys.stderr)


if __name__ == "__main__":
    main()
